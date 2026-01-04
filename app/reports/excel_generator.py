"""
Excel Report Generator for Microfinance MIS
Generates Excel reports for clients, loans, repayments, and BoG compliance
"""
from decimal import Decimal
from datetime import date, datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from clients.models import Client
from loans.models import Loan, LoanProduct, LoanSchedule
from repayments.models import Repayment
from dashboard.analytics import PortfolioAnalytics


class ExcelReportGenerator:
    """Generate professional Excel reports for microfinance operations."""
    
    def __init__(self):
        self.wb = None
        self.analytics = PortfolioAnalytics()
    
    def _create_workbook(self):
        """Create a new workbook with styling."""
        self.wb = Workbook()
        return self.wb
    
    def _style_header_row(self, ws, row=1):
        """Apply styling to header row."""
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        
        for cell in ws[row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    def _auto_size_columns(self, ws):
        """Auto-size all columns based on content."""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def generate_client_portfolio_report(self):
        """Generate comprehensive client portfolio report."""
        self._create_workbook()
        ws = self.wb.active
        ws.title = "Client Portfolio"
        
        # Headers
        headers = [
            'Client ID', 'Full Name', 'Phone', 'Email', 'Type', 
            'Risk Category', 'Status', 'KYC Verified', 'Ghana Card',
            'Active Loans', 'Total Outstanding (GHS)', 'Created Date'
        ]
        ws.append(headers)
        self._style_header_row(ws)
        
        # Data
        clients = Client.objects.all().order_by('-created_at')
        
        for client in clients:
            ws.append([
                client.client_id,
                client.full_name,
                client.phone,
                client.email,
                client.get_client_type_display(),
                client.get_risk_category_display(),
                client.get_status_display(),
                'Yes' if client.kyc_verified else 'No',
                client.ghana_card_id or 'N/A',
                client.get_active_loans_count(),
                float(client.get_total_outstanding()),
                client.created_at.strftime('%Y-%m-%d'),
            ])
        
        # Summary
        ws.append([])
        summary_row = ws.max_row + 1
        ws[f'A{summary_row}'] = 'SUMMARY'
        ws[f'A{summary_row}'].font = Font(bold=True, size=12)
        
        ws.append(['Total Clients:', clients.count()])
        ws.append(['Active Clients:', clients.filter(status='active').count()])
        ws.append(['KYC Verified:', clients.filter(kyc_verified=True).count()])
        
        self._auto_size_columns(ws)
        
        return self._save_workbook()
    
    def generate_loan_aging_report(self):
        """Generate loan aging report showing arrears."""
        self._create_workbook()
        ws = self.wb.active
        ws.title = "Loan Aging Report"
        
        # Headers
        headers = [
            'Loan ID', 'Client', 'Product', 'Principal (GHS)', 
            'Outstanding (GHS)', 'Disbursed Date', 'Maturity Date',
            'Days in Arrears', 'BoG Classification', 'Status'
        ]
        ws.append(headers)
        self._style_header_row(ws)
        
        # Data - Active loans ordered by days in arrears (descending)
        loans = Loan.objects.filter(status='active').order_by('-days_in_arrears')
        
        for loan in loans:
            ws.append([
                loan.loan_id,
                loan.client.full_name,
                loan.product.name if loan.product else 'N/A',
                float(loan.principal),
                float(loan.get_outstanding_balance()),
                loan.disbursed_date.strftime('%Y-%m-%d') if loan.disbursed_date else 'N/A',
                loan.maturity_date.strftime('%Y-%m-%d') if loan.maturity_date else 'N/A',
                loan.days_in_arrears,
                loan.get_classification_display(),
                loan.get_status_display(),
            ])
        
        # Aging buckets summary
        ws.append([])
        ws.append(['AGING ANALYSIS'])
        ws[f'A{ws.max_row}'].font = Font(bold=True, size=12)
        
        aging = self.analytics.get_arrears_aging()
        ws.append(['Aging Bucket', 'Loan Count', 'Amount (GHS)'])
        self._style_header_row(ws, ws.max_row)
        
        for bucket, data in aging.items():
            ws.append([bucket, data['count'], float(data['value'])])
        
        self._auto_size_columns(ws)
        
        return self._save_workbook()
    
    def generate_collection_report(self, start_date=None, end_date=None):
        """Generate collection performance report."""
        self._create_workbook()
        ws = self.wb.active
        ws.title = "Collection Report"
        
        if not start_date:
            start_date = date.today().replace(day=1)  # First day of month
        if not end_date:
            end_date = date.today()
        
        # Title
        ws['A1'] = f'Collection Report: {start_date} to {end_date}'
        ws['A1'].font = Font(bold=True, size=14)
        ws.append([])
        
        # Headers
        headers = [
            'Receipt #', 'Loan ID', 'Client', 'Amount (GHS)',
            'Date Paid', 'Method', 'Principal', 'Interest',
            'Penalty', 'Status', 'Received By'
        ]
        ws.append(headers)
        self._style_header_row(ws, 3)
        
        # Data
        repayments = Repayment.objects.filter(
            paid_on__gte=start_date,
            paid_on__lte=end_date
        ).select_related('loan', 'loan__client', 'received_by').order_by('-paid_on')
        
        total_collected = Decimal('0.00')
        
        for repayment in repayments:
            ws.append([
                repayment.receipt_number,
                repayment.loan.loan_id,
                repayment.loan.client.full_name,
                float(repayment.amount),
                repayment.paid_on.strftime('%Y-%m-%d'),
                repayment.get_method_display(),
                float(repayment.principal_paid),
                float(repayment.interest_paid),
                float(repayment.penalty_paid),
                repayment.get_status_display(),
                repayment.received_by.username if repayment.received_by else 'N/A',
            ])
            
            if repayment.status == 'confirmed':
                total_collected += repayment.amount
        
        # Summary
        ws.append([])
        summary_row = ws.max_row + 1
        ws[f'A{summary_row}'] = 'COLLECTION SUMMARY'
        ws[f'A{summary_row}'].font = Font(bold=True, size=12)
        
        ws.append(['Total Payments:', repayments.count()])
        ws.append(['Total Amount Collected (GHS):', float(total_collected)])
        ws.append(['Average Payment (GHS):', float(total_collected / repayments.count()) if repayments.count() > 0 else 0])
        
        self._auto_size_columns(ws)
        
        return self._save_workbook()
    
    def generate_bog_regulatory_report(self):
        """Generate Bank of Ghana regulatory compliance report."""
        self._create_workbook()
        ws = self.wb.active
        ws.title = "BoG Regulatory Report"
        
        # Title
        ws['A1'] = f'Bank of Ghana Prudential Report - {date.today().strftime("%B %Y")}'
        ws['A1'].font = Font(bold=True, size=14)
        ws.append([])
        
        # Portfolio Summary
        ws.append(['PORTFOLIO SUMMARY'])
        ws[f'A{ws.max_row}'].font = Font(bold=True, size=12)
        
        portfolio = self.analytics.get_portfolio_summary()
        
        ws.append(['Total Portfolio Value (GHS):', float(portfolio['total_portfolio_value'])])
        ws.append(['Total Outstanding (GHS):', float(portfolio['total_outstanding'])])
        ws.append(['Number of Active Loans:', portfolio['active_loans_count']])
        ws.append([])
        
        # Loan Classification (BoG Standard)
        ws.append(['LOAN CLASSIFICATION (BoG STANDARD)'])
        ws[f'A{ws.max_row}'].font = Font(bold=True, size=12)
        
        headers = ['Classification', 'Days in Arrears', 'Loan Count', 'Value (GHS)', 'Provision %', 'Provision Required (GHS)']
        ws.append(headers)
        self._style_header_row(ws, ws.max_row)
        
        classification = self.analytics.get_bog_classification()
        provision_data = self.analytics.get_provision_requirements()
        
        classification_labels = {
            'current': 'Current',
            'substandard': 'Substandard',
            'doubtful': 'Doubtful',
            'loss': 'Loss',
        }
        
        days_ranges = {
            'current': '0-30',
            'substandard': '31-90',
            'doubtful': '91-180',
            'loss': '180+',
        }
        
        provision_rates = {
            'current': 1,
            'substandard': 10,
            'doubtful': 50,
            'loss': 100,
        }
        
        for key, label in classification_labels.items():
            if key in classification:
                value = float(classification[key]['value'])
                provision_required = value * (provision_rates[key] / 100)
                
                ws.append([
                    label,
                    days_ranges[key],
                    classification[key]['count'],
                    value,
                    provision_rates[key],
                    provision_required,
                ])
        
        ws.append([])
        ws.append(['TOTAL PROVISIONS REQUIRED (GHS):', float(provision_data['total_provisions'])])
        
        # PAR Metrics
        ws.append([])
        ws.append(['PORTFOLIO AT RISK (PAR) METRICS'])
        ws[f'A{ws.max_row}'].font = Font(bold=True, size=12)
        
        par = self.analytics.get_par_metrics()
        
        ws.append(['PAR 30 (%):', f"{par['par_30_rate']:.2f}"])
        ws.append(['PAR 30 Value (GHS):', float(par['par_30_value'])])
        ws.append(['PAR 90 (%):', f"{par['par_90_rate']:.2f}"])
        ws.append(['PAR 90 Value (GHS):', float(par['par_90_value'])])
        
        # Product Performance
        ws.append([])
        ws.append(['PRODUCT PERFORMANCE'])
        ws[f'A{ws.max_row}'].font = Font(bold=True, size=12)
        
        headers = ['Product Name', 'Active Loans', 'Outstanding (GHS)', 'Arrears Loans', 'Arrears Value (GHS)']
        ws.append(headers)
        self._style_header_row(ws, ws.max_row)
        
        products = self.analytics.get_loan_product_performance()
        
        for product_data in products:
            ws.append([
                product_data['product'],
                product_data['active_loans'],
                float(product_data['outstanding']),
                product_data['arrears_loans'],
                float(product_data['arrears_value']),
            ])
        
        self._auto_size_columns(ws)
        
        return self._save_workbook()
    
    def generate_loan_schedule_report(self, loan_id):
        """Generate detailed repayment schedule for a specific loan."""
        try:
            loan = Loan.objects.get(loan_id=loan_id)
        except Loan.DoesNotExist:
            return None
        
        self._create_workbook()
        ws = self.wb.active
        ws.title = f"Schedule {loan_id}"
        
        # Loan details
        ws['A1'] = f'LOAN REPAYMENT SCHEDULE'
        ws['A1'].font = Font(bold=True, size=14)
        ws.append([])
        
        ws.append(['Loan ID:', loan.loan_id])
        ws.append(['Client:', loan.client.full_name])
        ws.append(['Product:', loan.product.name if loan.product else 'N/A'])
        ws.append(['Principal (GHS):', float(loan.principal)])
        ws.append(['Interest Rate (%):', float(loan.interest_rate)])
        ws.append(['Term (months):', loan.term_months])
        ws.append(['Disbursed Date:', loan.disbursed_date.strftime('%Y-%m-%d') if loan.disbursed_date else 'N/A'])
        ws.append(['Maturity Date:', loan.maturity_date.strftime('%Y-%m-%d') if loan.maturity_date else 'N/A'])
        ws.append([])
        
        # Schedule
        ws.append(['REPAYMENT SCHEDULE'])
        ws[f'A{ws.max_row}'].font = Font(bold=True, size=12)
        
        headers = [
            '#', 'Due Date', 'Principal Due', 'Interest Due',
            'Total Due', 'Principal Paid', 'Interest Paid', 
            'Total Paid', 'Balance', 'Status'
        ]
        ws.append(headers)
        self._style_header_row(ws, ws.max_row)
        
        schedule_entries = loan.schedule_entries.all().order_by('installment_number')
        
        for entry in schedule_entries:
            ws.append([
                entry.installment_number,
                entry.due_date.strftime('%Y-%m-%d'),
                float(entry.principal_due),
                float(entry.interest_due),
                float(entry.total_due),
                float(entry.principal_paid),
                float(entry.interest_paid),
                float(entry.total_paid),
                float(entry.balance_due),
                'Paid' if entry.is_paid else 'Pending',
            ])
        
        self._auto_size_columns(ws)
        
        return self._save_workbook()
    
    def _save_workbook(self):
        """Save workbook to BytesIO and return."""
        output = BytesIO()
        self.wb.save(output)
        output.seek(0)
        return output
